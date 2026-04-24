"""
藤田さん上長審査用の一覧表を生成。

出力:
    output/review_list_YYYYMMDD_HHMMSS.xlsx

項目:
    - 会社ID / 会社名 / マスタ状態
    - 建築許可の種類（業種一覧）
    - 行政庁 / 般特 / 許可年次 / 許可番号 / 許可日 / 有効期限
    - 必要8書類 揃い判定（各書類○×と集計）
    - 最終受信日
"""
from __future__ import annotations

import sqlite3
import sys
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DB = PROJECT_ROOT / "data" / "permit_tracker.db"
OUTPUT = PROJECT_ROOT / "output"
MASTER_XLSX = Path(r"C:/Users/owner/Desktop/145社_許可一覧.xlsx")

REQUIRED_DOCS = [
    "取引申請書", "建設業許可証", "決算書", "工事経歴書",
    "取引先一覧表", "労働安全衛生誓約書", "資格略字一覧", "労働者名簿",
]


def load_master_names() -> set[str]:
    wb = openpyxl.load_workbook(MASTER_XLSX, data_only=True)
    ws = wb["145社許可一覧"]
    names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            names.add(str(row[0]).strip())
    return names


def main():
    conn = sqlite3.connect(str(DB))
    conn.row_factory = sqlite3.Row
    master = load_master_names()

    # 全companies
    all_comps = conn.execute(
        "SELECT company_id, official_name, status FROM companies ORDER BY company_id"
    ).fetchall()

    # 各会社のdata集計
    rows_for_sheet = []
    for comp in all_comps:
        cid = comp["company_id"]
        name = comp["official_name"]

        # マスタ状態
        in_master = "✓" if name in master else "×"

        # 最新許可（current_flag=1）
        permit = conn.execute(
            "SELECT p.permit_number, p.permit_authority, p.permit_category, "
            "p.permit_year, p.issue_date, p.expiry_date, "
            "GROUP_CONCAT(pt.trade_name, '、') AS trades "
            "FROM permits p LEFT JOIN permit_trades pt ON pt.permit_id = p.permit_id "
            "WHERE p.company_id=? AND p.current_flag=1 "
            "GROUP BY p.permit_id "
            "ORDER BY p.expiry_date DESC LIMIT 1",
            (cid,),
        ).fetchone()

        # 書類揃い判定
        doc_types_present = {r["doc_type_name"] for r in conn.execute(
            "SELECT DISTINCT doc_type_name FROM pages WHERE company_id=?", (cid,)
        )}
        doc_checks = {dt: ("○" if dt in doc_types_present else "×") for dt in REQUIRED_DOCS}
        has_count = sum(1 for dt in REQUIRED_DOCS if dt in doc_types_present)
        all_ok = "◎" if has_count == len(REQUIRED_DOCS) else ("△" if has_count > 0 else "")

        # 最新受信日（inbound_messages）
        last_recv = conn.execute(
            "SELECT MAX(received_at) FROM inbound_messages WHERE company_id=?", (cid,)
        ).fetchone()[0]

        rows_for_sheet.append({
            "会社ID": cid,
            "会社名": name,
            "マスタ": in_master,
            "状態": comp["status"],
            "行政庁": permit["permit_authority"] if permit else "",
            "般特": permit["permit_category"] if permit else "",
            "許可年次": permit["permit_year"] if permit else "",
            "許可番号": permit["permit_number"] if permit else "",
            "許可日": permit["issue_date"] if permit else "",
            "有効期限": permit["expiry_date"] if permit else "",
            "建築許可の種類": permit["trades"] if permit and permit["trades"] else "",
            **doc_checks,
            "揃い": f"{has_count}/{len(REQUIRED_DOCS)}",
            "審査可": all_ok,
            "最終受信": last_recv or "",
        })

    # Excel 出力
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "審査一覧"

    headers = [
        "会社ID", "会社名", "マスタ", "状態",
        "行政庁", "般特", "許可年次", "許可番号", "許可日", "有効期限", "建築許可の種類",
        *REQUIRED_DOCS,
        "揃い", "審査可", "最終受信",
    ]
    ws.append(headers)

    # ヘッダ書式
    header_fill = PatternFill(start_color="1B3D6F", end_color="1B3D6F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # データ行
    for rec in rows_for_sheet:
        ws.append([rec.get(h, "") for h in headers])

    # 列幅
    widths = {
        "会社ID": 9, "会社名": 28, "マスタ": 7, "状態": 10,
        "行政庁": 14, "般特": 6, "許可年次": 8, "許可番号": 36,
        "許可日": 12, "有効期限": 12, "建築許可の種類": 40,
        "揃い": 7, "審査可": 7, "最終受信": 20,
    }
    for doc in REQUIRED_DOCS:
        widths[doc] = 10
    for i, h in enumerate(headers, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = widths.get(h, 12)

    # フィルタ
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "D2"

    # 審査可セル色付け
    green = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    yellow = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    approve_col = headers.index("審査可") + 1
    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, approve_col).value
        if val == "◎":
            ws.cell(r, approve_col).fill = green
        elif val == "△":
            ws.cell(r, approve_col).fill = yellow

    # 出力
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUTPUT / f"review_list_{ts}.xlsx"
    wb.save(out_path)
    print(f"出力: {out_path}")
    print(f"総件数: {len(rows_for_sheet)}")
    # 集計
    ready = sum(1 for r in rows_for_sheet if r["審査可"] == "◎")
    partial = sum(1 for r in rows_for_sheet if r["審査可"] == "△")
    none = sum(1 for r in rows_for_sheet if not r["審査可"])
    print(f"  ◎全書類揃い(審査可): {ready}社")
    print(f"  △一部受信: {partial}社")
    print(f"  （書類なし）: {none}社")


if __name__ == "__main__":
    main()
