"""
Phase R-3: 担当者編集ワークブックの反映スクリプト。

設計原則 (GPT-5.5 反映):
- 原本不変: v2 一覧表は直接書き換えず、新版 v3 を生成
- 二重反映防止: workflow_id を decisions_log で一意管理。再 apply は拒否
- staging 経由: 新規候補は companies_staging のみ、本体 companies に直接 INSERT 禁止
- transaction: DB 更新は単一 transaction、失敗時 ROLLBACK
- dry-run デフォルト: 何が変わるかを CSV 出力、DB / ファイル変更なし
- snapshot 復元: --rollback で Phase R-0 のバックアップから復元

Usage:
  python scripts/apply_editor_decisions.py <workbook_path>           # dry-run
  python scripts/apply_editor_decisions.py <workbook_path> --execute # 実行
  python scripts/apply_editor_decisions.py --rollback <TS>           # snapshot 復元
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from shutil import copy2

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

PROJECT = Path(__file__).resolve().parent.parent
DB_PATH = PROJECT / "data" / "permit_tracker.db"
V2_XLSX = PROJECT / "output" / "FDE_MANAGED" / "00_一覧表_v2_20260427.xlsx"
SNAP_DIR = PROJECT / "data" / "snapshots"
OUT_SNAP_DIR = PROJECT / "output" / "snapshots"

DOC_COLS = ["取引申請書", "建設業許可証", "決算書", "工事経歴書",
             "取引先一覧表", "労働安全衛生誓約書", "資格略字一覧", "労働者名簿"]


def hash_decision(*parts: str) -> str:
    s = "|".join(str(p) if p is not None else "" for p in parts)
    return hashlib.sha256(s.encode("utf-8")).hexdigest()[:16]


def load_meta(wb) -> dict:
    ws = wb["0_workflow_meta"]
    meta = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        v = ws.cell(r, 2).value
        if k:
            meta[k] = v
    return meta


def read_sheet(wb, name: str, header_row=1) -> list[dict]:
    ws = wb[name]
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        # Skip empty rows
        if all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1)):
            continue
        rows.append({h: ws.cell(r, c).value for c, h in enumerate(headers, 1) if h})
    return rows


def collect_decisions(wb) -> dict:
    """各シートから decision を抽出"""
    meta = load_meta(wb)
    return {
        "meta": meta,
        "sheet1_mapping": read_sheet(wb, "1_突合せ確認"),
        "sheet2_docs": read_sheet(wb, "2_受領状況編集"),
        "sheet3_expiry": read_sheet(wb, "3_期限管理"),
        "sheet4_new": read_sheet(wb, "4_新規候補"),
        "sheet5_inactive": read_sheet(wb, "5_既存対象外化候補"),
    }


def precheck(decisions: dict, conn: sqlite3.Connection) -> list[str]:
    """事前検証。失敗を一覧で返す (空なら OK)"""
    fails = []
    meta = decisions["meta"]

    # 1. workflow_id 存在
    wf = meta.get("workflow_id")
    if not wf:
        fails.append("workflow_id が meta シートにない")
        return fails

    # 2. 二重反映チェック
    n = conn.execute(
        "SELECT COUNT(*) FROM decisions_log WHERE workflow_id=?", (wf,)
    ).fetchone()[0]
    if n > 0:
        fails.append(f"workflow_id={wf} は既に {n} 件の decision が反映済み (二重反映防止)")

    # 3. v2 ファイル一致
    v2_in_meta = str(meta.get("対象 v2 ファイル", "")).replace("\\", "/")
    v2_actual = str(V2_XLSX.relative_to(PROJECT)).replace("\\", "/")
    if v2_in_meta and v2_in_meta not in (v2_actual, str(V2_XLSX)):
        fails.append(f"対象 v2 が異なる: meta={v2_in_meta}, actual={v2_actual}")

    # 4. expected_decision_count vs 実数
    expected = meta.get("expected_decision_count")
    actual = sum(len(decisions[k]) for k in decisions if k.startswith("sheet"))
    if expected and int(expected) != actual:
        fails.append(f"decision 数が meta と不一致: expected={expected}, actual={actual}")

    return fails


def apply_decisions(decisions: dict, conn: sqlite3.Connection,
                     execute: bool, ts: str) -> dict:
    """反映処理。dry-run なら DB 変更なし"""
    meta = decisions["meta"]
    workflow_id = meta.get("workflow_id")
    operator = "fde_apply_script"

    summary = {
        "doc_status_apply": 0,
        "doc_status_hold": 0,
        "doc_status_reject": 0,
        "expiry_status_apply": 0,
        "staging_add": 0,
        "inactivate": 0,
        "skipped_no_decision": 0,
    }
    log_rows = []
    v2_updates = []  # (company_id, doc_col, new_value)

    # Sheet 2: 書類状態反映
    for d in decisions["sheet2_docs"]:
        judge = d.get("担当者判断 [反映/保留/却下]") or ""
        if judge == "反映":
            new_val = d.get("反映後の値 [○/×/対象外/-]") or ""
            v2_updates.append({
                "company_id": d.get("company_id"),
                "doc_col": d.get("書類カテゴリ"),
                "new_value": new_val,
                "decision_id": d.get("decision_id"),
            })
            log_rows.append({
                "decision_id": d.get("decision_id"),
                "workflow_id": workflow_id,
                "company_id": d.get("company_id"),
                "document_type": d.get("書類カテゴリ"),
                "decision_kind": "doc_status",
                "old_value": d.get("当方v2の値"),
                "new_value": new_val,
                "decision_value": "反映",
                "operator": operator,
                "fujita_comment": d.get("藤田指摘原文"),
                "operator_comment": d.get("コメント (自由記入)"),
                "source_workbook": meta.get("workflow_id", ""),
            })
            summary["doc_status_apply"] += 1
        elif judge == "保留":
            summary["doc_status_hold"] += 1
        elif judge == "却下":
            summary["doc_status_reject"] += 1
        else:
            summary["skipped_no_decision"] += 1

    # Sheet 3: 期限ステータス
    for d in decisions["sheet3_expiry"]:
        st = d.get("期限ステータス [MLIT更新確認済/藤田に再依頼/期限切れ確定/督促済]") or ""
        if not st:
            continue
        log_rows.append({
            "decision_id": d.get("decision_id"),
            "workflow_id": workflow_id,
            "company_id": d.get("company_id"),
            "document_type": None,
            "decision_kind": "expiry_status",
            "old_value": d.get("藤田判定"),
            "new_value": st,
            "decision_value": st,
            "operator": operator,
            "fujita_comment": None,
            "operator_comment": d.get("コメント"),
            "source_workbook": meta.get("workflow_id", ""),
        })
        summary["expiry_status_apply"] += 1

    # Sheet 4: 新規候補 -> staging
    for d in decisions["sheet4_new"]:
        proc = d.get("処理 [候補確定/別名疑い/対象外候補]") or ""
        if proc != "候補確定":
            continue
        log_rows.append({
            "decision_id": d.get("decision_id"),
            "workflow_id": workflow_id,
            "company_id": d.get("提案company_id"),
            "document_type": None,
            "decision_kind": "staging_add",
            "old_value": None,
            "new_value": d.get("fujita_name"),
            "decision_value": "候補確定",
            "operator": operator,
            "fujita_comment": None,
            "operator_comment": d.get("コメント"),
            "source_workbook": meta.get("workflow_id", ""),
        })
        if execute:
            conn.execute(
                "INSERT INTO companies_staging (fujita_row_id, fujita_name, proposed_company_id, workflow_id, notes) VALUES (?, ?, ?, ?, ?)",
                (d.get("fujita_row_id"), d.get("fujita_name"),
                 d.get("提案company_id"), workflow_id, d.get("コメント")),
            )
        summary["staging_add"] += 1

    # Sheet 5: INACTIVE 化
    for d in decisions["sheet5_inactive"]:
        proc = d.get("処理 [INACTIVE化/対象外フラグ/維持]") or ""
        if proc not in ("INACTIVE化", "対象外フラグ"):
            continue
        log_rows.append({
            "decision_id": d.get("decision_id"),
            "workflow_id": workflow_id,
            "company_id": d.get("company_id"),
            "document_type": None,
            "decision_kind": "inactivate",
            "old_value": d.get("現状態"),
            "new_value": proc,
            "decision_value": proc,
            "operator": operator,
            "fujita_comment": d.get("藤田判定理由 (memo)"),
            "operator_comment": d.get("コメント"),
            "source_workbook": meta.get("workflow_id", ""),
        })
        if execute:
            new_state = "INACTIVE" if proc == "INACTIVE化" else "MANAGED_OUT"
            conn.execute(
                "UPDATE companies SET status=?, updated_at=datetime('now','localtime') WHERE company_id=?",
                (new_state, d.get("company_id")),
            )
        summary["inactivate"] += 1

    # decisions_log INSERT
    if execute:
        for lr in log_rows:
            conn.execute(
                """INSERT INTO decisions_log
                (decision_id, workflow_id, company_id, document_type, decision_kind,
                 old_value, new_value, decision_value, operator, fujita_comment,
                 operator_comment, source_workbook)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (lr["decision_id"], lr["workflow_id"], lr["company_id"],
                 lr["document_type"], lr["decision_kind"], lr["old_value"],
                 lr["new_value"], lr["decision_value"], lr["operator"],
                 lr["fujita_comment"], lr["operator_comment"],
                 lr["source_workbook"]),
            )

    return {"summary": summary, "log_rows": log_rows, "v2_updates": v2_updates}


def generate_v3(v2_updates: list[dict], ts: str, workflow_id: str) -> Path:
    """v2 を読み、v2_updates を反映した新ファイル v3 を生成 (原本不変)"""
    wb = openpyxl.load_workbook(V2_XLSX)
    ws = wb["マスタ145社受領状況"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    cid_col = headers.index("会社ID") + 1
    doc_col_idx = {dc: headers.index(dc) + 1 for dc in DOC_COLS if dc in headers}

    # 行 index
    cid_to_row = {}
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, cid_col).value
        if cid:
            cid_to_row[cid] = r

    n_applied = 0
    for u in v2_updates:
        r = cid_to_row.get(u["company_id"])
        c = doc_col_idx.get(u["doc_col"])
        if r and c:
            ws.cell(r, c).value = u["new_value"]
            n_applied += 1

    out = PROJECT / "output" / "FDE_MANAGED" / f"00_一覧表_v3_{ts}_{workflow_id[:8]}.xlsx"
    wb.save(out)
    print(f"  v3 生成: {out.relative_to(PROJECT)}  反映 {n_applied}/{len(v2_updates)} セル")
    return out


def write_dryrun_csv(result: dict, ts: str) -> Path:
    out = PROJECT / "output" / "FDE_MANAGED" / f"phase_r3_dryrun_{ts}.csv"
    with out.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["decision_id", "workflow_id", "company_id", "document_type",
                    "decision_kind", "old_value", "new_value", "decision_value",
                    "fujita_comment", "operator_comment"])
        for lr in result["log_rows"]:
            w.writerow([lr["decision_id"], lr["workflow_id"], lr["company_id"],
                        lr["document_type"], lr["decision_kind"], lr["old_value"],
                        lr["new_value"], lr["decision_value"],
                        str(lr["fujita_comment"] or "")[:200],
                        str(lr["operator_comment"] or "")[:200]])
    print(f"  dry-run CSV: {out.relative_to(PROJECT)}")
    return out


def rollback(ts: str) -> None:
    """snapshot 復元方式の rollback"""
    print(f"=== Rollback to TS={ts} ===")
    manifest_path = SNAP_DIR / f"manifest_{ts}.json"
    if not manifest_path.exists():
        print(f"ERROR: manifest not found: {manifest_path}")
        sys.exit(1)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for item in manifest["items"]:
        src = PROJECT / item["dst"]
        dst = PROJECT / item["src"]
        if not src.exists():
            print(f"  [WARN] missing snapshot: {src}")
            continue
        copy2(src, dst)
        print(f"  restored: {dst.relative_to(PROJECT)} <- {src.relative_to(PROJECT)}")
    print()
    print("注意:")
    print("  - GAS スプレッドシートは Sheets UI のバージョン履歴で手動復元してください")
    print(f"  - decisions_log に Phase R 反映済 workflow_id があれば手動で確認してください")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", nargs="?", type=Path)
    ap.add_argument("--execute", action="store_true")
    ap.add_argument("--rollback", help="snapshot TS を指定して復元")
    args = ap.parse_args()

    if args.rollback:
        rollback(args.rollback)
        return

    if not args.workbook or not args.workbook.exists():
        print("ERROR: workbook を指定してください", file=sys.stderr)
        sys.exit(1)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    print(f"=== Phase R-3 apply_editor_decisions ===")
    print(f"  TS: {ts}")
    print(f"  workbook: {args.workbook.relative_to(PROJECT) if args.workbook.is_relative_to(PROJECT) else args.workbook}")
    print(f"  mode: {'EXECUTE' if args.execute else 'DRY-RUN'}")

    print("\nLoading workbook...")
    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    decisions = collect_decisions(wb)
    print(f"  meta: workflow_id={decisions['meta'].get('workflow_id')}")
    for k, v in decisions.items():
        if k != "meta":
            print(f"  {k}: {len(v)} rows")

    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")

    print("\n=== Pre-check ===")
    fails = precheck(decisions, conn)
    if fails:
        print("  FAILED:")
        for f in fails:
            print(f"    - {f}")
        if args.execute:
            print("\nApply 中止 (precheck failed)")
            sys.exit(2)
        else:
            print("\nDRY-RUN なので続行")
    else:
        print("  OK")

    print("\n=== Apply ===")
    if args.execute:
        try:
            conn.execute("BEGIN")
            result = apply_decisions(decisions, conn, execute=True, ts=ts)
            conn.commit()
        except Exception as e:
            conn.rollback()
            print(f"  EXCEPTION: {e}", file=sys.stderr)
            raise
    else:
        result = apply_decisions(decisions, conn, execute=False, ts=ts)

    print("  Summary:")
    for k, v in result["summary"].items():
        print(f"    {k}: {v}")

    if result["v2_updates"]:
        print("\n=== v3 generation ===")
        if args.execute:
            generate_v3(result["v2_updates"], ts, decisions["meta"]["workflow_id"])
        else:
            print(f"  would update {len(result['v2_updates'])} cells (dry-run, skip)")

    print("\n=== Output ===")
    write_dryrun_csv(result, ts)

    conn.close()
    print()
    print("Done.")


if __name__ == "__main__":
    main()
