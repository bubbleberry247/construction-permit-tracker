"""
一覧表 v2 の有効期限は添付許可証 (旧期限) のまま、備考で更新済を明示。

対象: 4-27 強制更新で許可証が更新済と判明した社 (C0071, C0117)
動作:
  1. MLITPermits シートから最新の許可情報 (新期限) 取得
  2. 一覧表 v2 の該当行の「有効期限」列を旧期限 (添付値) に戻す
  3. 備考列に「添付許可証の期限は旧期限 YYYY-MM-DD だが、MLIT 公式で更新確認済。
     新期限は YYYY-MM-DD」を記述
  4. 既存 COMPANY_NOTES (C0008 三和シャッター 等) も保持

Usage:
  python scripts/patch_renewed_permits_in_review.py <xlsx_path>
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl
import gspread
from google.oauth2.service_account import Credentials

sys.stdout.reconfigure(encoding="utf-8")

PROJECT = Path(__file__).resolve().parent.parent
CONFIG = PROJECT / "config.json"

# 許可証が更新されていた社 (期限が伸びた、Stage 4 強制更新で判明)
RENEWED_TARGETS = {
    "C0071": {"old_expiry": "2026-04-13", "company_name": "株式会社谷野宮組"},
    "C0117": {"old_expiry": "2026-04-06", "company_name": "株式会社三富"},
}


def open_sheet():
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    creds = Credentials.from_service_account_file(
        cfg["GOOGLE_SERVICE_ACCOUNT_FILE"],
        scopes=["https://www.googleapis.com/auth/spreadsheets",
                "https://www.googleapis.com/auth/drive"],
    )
    gc = gspread.authorize(creds)
    return gc.open_by_key(cfg["GOOGLE_SHEETS_ID"])


def fetch_mlit_data(sh, company_id: str) -> dict | None:
    ws = sh.worksheet("MLITPermits")
    rows = ws.get_all_values()
    hdr = rows[0]
    cid_i = hdr.index("company_id")
    for r in rows[1:]:
        if len(r) > cid_i and r[cid_i] == company_id:
            return {h: r[i] if i < len(r) else "" for i, h in enumerate(hdr)}
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx_path", type=Path)
    args = ap.parse_args()

    if not args.xlsx_path.exists():
        print(f"ERROR: {args.xlsx_path} not found", file=sys.stderr)
        sys.exit(1)

    sh = open_sheet()

    # MLITPermits から各 target の最新データ取得
    mlit_data = {}
    for cid in RENEWED_TARGETS:
        d = fetch_mlit_data(sh, cid)
        if d:
            mlit_data[cid] = d
            print(f"MLIT 取得: {cid} {d.get('company_name')} 新期限={d.get('expiry_date')} 残{d.get('days_remaining')}日")
        else:
            print(f"  [WARN] {cid} は MLITPermits に存在しない")

    # 一覧表 v2 を編集
    wb = openpyxl.load_workbook(args.xlsx_path)
    ws = wb["マスタ145社受領状況"]
    hdr = [c.value for c in ws[1]]
    cid_col = hdr.index("会社ID") + 1
    note_col = hdr.index("備考") + 1
    exp_col = hdr.index("有効期限") + 1
    issue_col = hdr.index("許可日") + 1
    cat_col = hdr.index("般特") + 1
    year_col = hdr.index("許可年次") + 1
    pnum_col = hdr.index("許可番号") + 1
    auth_col = hdr.index("行政庁") + 1

    n_updated = 0
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row, cid_col).value
        if cid not in RENEWED_TARGETS:
            continue
        target_info = RENEWED_TARGETS[cid]
        mlit = mlit_data.get(cid)
        if not mlit:
            continue

        # 現在の有効期限 (前回 patch で MLIT 新値に上書きされている可能性あり)
        current_expiry_in_sheet = ws.cell(row, exp_col).value or "(未設定)"

        # 一覧表の有効期限は「添付許可証の期限」(旧期限) に戻す
        old_expiry = target_info["old_expiry"]
        new_expiry = mlit.get("expiry_date", "")

        ws.cell(row, exp_col).value = old_expiry

        # 備考: 添付許可証期限と MLIT で確認した新期限を明示
        old_note = ws.cell(row, note_col).value or ""
        new_note_text = (
            f"【許可証更新済】添付許可証の有効期限 ({old_expiry}) は更新前のもの。"
            f"MLIT 公式で確認したところ更新済。新有効期限: {new_expiry} (5年延長)。"
        )

        # 既存 note から前回 patch のテキストがあれば除去 (再実行対応)
        if "添付書類は旧許可証時代" in old_note or "【許可証更新済】" in old_note:
            parts = [p.strip() for p in old_note.split(" / ")]
            parts = [p for p in parts if "添付書類は旧許可証時代" not in p
                                       and "【許可証更新済】" not in p]
            old_note = " / ".join(parts)

        if old_note:
            ws.cell(row, note_col).value = f"{old_note} / {new_note_text}"
        else:
            ws.cell(row, note_col).value = new_note_text

        print(f"\n更新: {cid} {target_info['company_name']}")
        print(f"  有効期限: {current_expiry_in_sheet} → {old_expiry} (添付許可証の値に復元)")
        print(f"  MLIT 新期限: {new_expiry} (備考に記載)")
        print(f"  備考: {ws.cell(row, note_col).value[:100]}...")
        n_updated += 1

    wb.save(args.xlsx_path)
    print(f"\n=== 完了 ===")
    print(f"  対象ファイル: {args.xlsx_path}")
    print(f"  更新行数: {n_updated}")


if __name__ == "__main__":
    main()
