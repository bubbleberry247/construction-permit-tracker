"""
納品用 Excel ダッシュボードを生成する。

入力:
  - data/permit_tracker.db
  - output/FDE_MANAGED/00_index/index.csv
  - config/required_docs.json

出力:
  output/FDE_MANAGED/dashboard.xlsx

シート構成（6 シート）:
  1. Dashboard    : ACTIVE 会社 × 9 書類マトリクス、hyperlink セル
  2. Documents    : 会社×書類×ファイル粒度の詳細一覧
  3. Candidates   : 要確認ページ
  4. Inactive     : status=INACTIVE 会社
  5. DeliveryLog  : 納品履歴（毎回追記、前回ブック継承）
  6. Config       : メタ情報

重要:
  リンクは =HYPERLINK() 数式ではなく openpyxl の cell.hyperlink 属性で設定する。
  数式は link_location に 255 文字制限があり、URI エンコードした日本語パスが
  この制限を超えると Excel が修復時に数式を削除してしまうため。
  cell.hyperlink は Excel 正規の XML <hyperlink> を使うため制限がなく、
  URI エンコード不要で日本語パスをそのまま扱える。

Usage:
  python scripts/generate_delivery_dashboard.py
  python scripts/generate_delivery_dashboard.py --package-id v1.0 --note "初期納品"
"""
from __future__ import annotations

import argparse
import csv
import json
import sqlite3
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parent.parent
DB = PROJECT / "data" / "permit_tracker.db"
CONFIG = PROJECT / "config" / "required_docs.json"
OUTPUT_ROOT = PROJECT / "output" / "FDE_MANAGED"
INDEX_CSV = OUTPUT_ROOT / "00_index" / "index.csv"
DASHBOARD_XLSX = OUTPUT_ROOT / "dashboard.xlsx"

REQUIRED_DOCS = json.loads(CONFIG.read_text(encoding="utf-8"))

HEADER_FILL = PatternFill(start_color="1B3D6F", end_color="1B3D6F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
GREEN_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
RED_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
GRAY_FILL = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
BORDER = Border(left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
                top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def link_target(rel_path: str, page_no: int | None = None) -> str:
    """openpyxl cell.hyperlink 用の相対パス文字列を返す。

    - スラッシュ区切り
    - URI エンコードしない（Excel 正規 hyperlink XML が日本語をそのまま扱える）
    - page_no があれば #page=N を付与（Acrobat/Edge 対応）
    """
    path = rel_path.replace("\\", "/")
    if page_no and page_no > 0:
        path += f"#page={page_no}"
    return path


def set_link_cell(cell, target: str, display: str):
    """セルに値と hyperlink を設定。Excel 数式ではなく正規 hyperlink XML を使う。"""
    cell.value = display
    cell.hyperlink = target
    cell.font = LINK_FONT


def load_index_csv() -> list[dict]:
    rows = []
    with INDEX_CSV.open("r", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            rows.append(r)
    return rows


def load_companies(conn: sqlite3.Connection) -> list[dict]:
    return [dict(r) for r in conn.execute("""
        SELECT c.company_id, c.official_name, c.corporation_type, c.status,
               CASE WHEN EXISTS (
                 SELECT 1 FROM permits p WHERE p.company_id = c.company_id AND p.current_flag = 1
               ) THEN 1 ELSE 0 END AS has_construction_permit
        FROM companies c ORDER BY c.company_id
    """)]


def load_candidates(conn: sqlite3.Connection) -> list[dict]:
    MANUAL = "'viewer_correction','manual_review_v2','manual_review','web_viewer','user_manual','manual_visual'"
    sql = f"""
      SELECT p.page_id, p.company_id, c.official_name, p.file_name, p.page_no,
             p.doc_type_name, p.confidence
      FROM pages p JOIN companies c ON c.company_id=p.company_id
      WHERE p.doc_type_name IN ('その他/不明', 'その他')
         OR (p.confidence IS NOT NULL AND p.confidence < 0.8)
         OR (p.confidence IS NULL AND p.company_id NOT IN
             (SELECT DISTINCT company_id FROM field_reviews WHERE confirmed_by IN ({MANUAL})))
      ORDER BY p.company_id, p.file_name, p.page_no
    """
    return [dict(r) for r in conn.execute(sql)]


def preserve_delivery_log(wb_old_path: Path | None) -> list[list]:
    if not wb_old_path or not wb_old_path.exists():
        return []
    try:
        wb = openpyxl.load_workbook(wb_old_path, read_only=True)
        if "DeliveryLog" not in wb.sheetnames:
            return []
        ws = wb["DeliveryLog"]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue
            if row and any(v for v in row):
                rows.append(list(row))
        return rows
    except Exception as e:
        print(f"WARN: 前回 DeliveryLog 読込失敗: {e}", file=sys.stderr)
        return []


def write_dashboard_sheet(ws, companies: list[dict], index_rows: list[dict]):
    headers = ["company_id", "会社名", "建設業"] + [d["display"] for d in REQUIRED_DOCS] + ["揃い率", "最終更新"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    received_by_company: dict[str, set[str]] = defaultdict(set)
    link_by_company_doctype: dict[tuple[str, str], tuple[str, int]] = {}
    for r in index_rows:
        cid = r["company_code"]
        dt_id = r["doc_type_id"]
        if dt_id in ("bundle", "unknown"):
            continue
        received_by_company[cid].add(dt_id)
        if (cid, dt_id) not in link_by_company_doctype:
            link_by_company_doctype[(cid, dt_id)] = (r["rel_path"], int(r.get("page_from") or 0))

    active = [c for c in companies if c["status"] == "ACTIVE"]
    # (sheet_row, col, rel_path, page_from) の hyperlink 設定タスクを記憶して後段で設定
    hyperlink_tasks: list[tuple[int, int, str, int]] = []

    for c in active:
        cid = c["company_id"]
        is_permit = bool(c["has_construction_permit"])
        row = [cid, c["official_name"], "○" if is_permit else "—"]
        required_count = 0
        received_count = 0
        col_idx_after_base = 4  # 会社×建設業=3列、書類はここから
        for i, d in enumerate(REQUIRED_DOCS):
            col_abs = col_idx_after_base + i
            applies = (d["applies_to"] == "all_active") or \
                      (d["applies_to"] == "has_construction_permit" and is_permit)
            if not applies:
                row.append("—")
                continue
            required_count += 1
            has = d["id"] in received_by_company[cid]
            if has:
                received_count += 1
                rel_path, page_from = link_by_company_doctype[(cid, d["id"])]
                row.append("○")
                # ws.append 後にセルを取得して hyperlink 設定するため、行番号は後で確定
                hyperlink_tasks.append((0, col_abs, rel_path, page_from))  # row=0 は placeholder
            else:
                row.append("×")
        rate = f"{received_count}/{required_count}" if required_count > 0 else "0/0"
        row.extend([rate, datetime.now().strftime("%Y-%m-%d")])
        ws.append(row)
        # 直前 append した行番号 = ws.max_row。このループ中の hyperlink_tasks の行番号を確定
        for j in range(len(hyperlink_tasks) - 1, -1, -1):
            if hyperlink_tasks[j][0] == 0:
                hyperlink_tasks[j] = (ws.max_row, hyperlink_tasks[j][1], hyperlink_tasks[j][2], hyperlink_tasks[j][3])
            else:
                break

    # hyperlink 設定
    for r, c, rel_path, page_from in hyperlink_tasks:
        cell = ws.cell(r, c)
        set_link_cell(cell, link_target(rel_path, page_from if page_from > 0 else None), "○")

    widths = [9, 30, 8] + [14] * len(REQUIRED_DOCS) + [9, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # セル装飾
    for r in range(2, ws.max_row + 1):
        for col in range(4, 4 + len(REQUIRED_DOCS)):
            cell = ws.cell(r, col)
            cell.alignment = CENTER
            cell.border = BORDER
            v = cell.value
            if v == "○" or v == "◎":
                cell.fill = GREEN_FILL
            elif v == "△":
                cell.fill = YELLOW_FILL
            elif v == "×":
                cell.fill = RED_FILL
            elif v == "—":
                cell.fill = GRAY_FILL

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions


def write_documents_sheet(ws, index_rows: list[dict]):
    headers = ["record_id", "company_code", "company_name", "doc_type",
               "original_filename", "rel_path", "page_from", "page_to",
               "received_date", "retention_until", "is_bundle", "link"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    link_col = headers.index("link") + 1
    for r in index_rows:
        page_from = int(r.get("page_from") or 0)
        ws.append([
            r["record_id"], r["company_code"], r["company_name"], r["doc_type"],
            r["original_filename"], r["rel_path"], page_from, int(r.get("page_to") or 0),
            r.get("received_date", ""), r.get("retention_until", ""), int(r.get("is_bundle", 0)),
            "",  # link セル値は後段で設定
        ])
        cell = ws.cell(ws.max_row, link_col)
        set_link_cell(cell, link_target(r["rel_path"], page_from if page_from > 0 else None), "開く")

    widths = [9, 10, 28, 18, 48, 60, 9, 9, 12, 12, 9, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions


def write_candidates_sheet(ws, candidates: list[dict]):
    headers = ["page_id", "company_id", "会社名", "file_name", "page_no", "doc_type_name", "confidence"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    for r in candidates:
        ws.append([r["page_id"], r["company_id"], r["official_name"], r["file_name"],
                   r["page_no"], r["doc_type_name"], r["confidence"]])

    widths = [9, 10, 28, 50, 7, 18, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_inactive_sheet(ws, companies: list[dict]):
    headers = ["company_id", "会社名", "corporation_type", "建設業"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    inactive = [c for c in companies if c["status"] == "INACTIVE"]
    for c in inactive:
        ws.append([c["company_id"], c["official_name"], c["corporation_type"] or "",
                   "○" if c["has_construction_permit"] else "—"])

    for i, w in enumerate([9, 30, 16, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_delivery_log_sheet(ws, previous_log: list[list], new_row: list):
    headers = ["delivered_at", "package_id", "type", "added_companies", "added_docs", "fde_name", "note"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    for row in previous_log:
        ws.append(row)
    if new_row:
        ws.append(new_row)

    for i, w in enumerate([20, 12, 8, 20, 12, 14, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_config_sheet(ws, companies: list[dict], index_rows: list[dict]):
    headers = ["key", "value"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    active_n = sum(1 for c in companies if c["status"] == "ACTIVE")
    permit_n = sum(1 for c in companies if c["has_construction_permit"])
    rows = [
        ["生成日時", datetime.now().isoformat()],
        ["companies 総数", len(companies)],
        ["ACTIVE 会社数", active_n],
        ["建設業許可保有", permit_n],
        ["index レコード数", len(index_rows)],
        ["required_docs バージョン", "1.0"],
        ["data_root", str(OUTPUT_ROOT)],
        ["FDE 連絡先", "kalimistk@gmail.com"],
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--package-id", default=f"v{datetime.now().strftime('%Y%m%d')}", help="納品パッケージ ID")
    ap.add_argument("--type", default="追加", choices=["初期", "追加"], help="納品種別")
    ap.add_argument("--note", default="", help="納品備考")
    ap.add_argument("--fde-name", default="FDE", help="FDE 氏名")
    args = ap.parse_args()

    if not INDEX_CSV.exists():
        print(f"ERROR: {INDEX_CSV} が無い。先に build_index_csv.py を実行してください。", file=sys.stderr)
        sys.exit(1)

    conn = sqlite3.connect(str(DB))
    conn.row_factory = sqlite3.Row

    companies = load_companies(conn)
    index_rows = load_index_csv()
    candidates = load_candidates(conn)
    previous_log = preserve_delivery_log(DASHBOARD_XLSX)

    print(f"会社: {len(companies)} (ACTIVE: {sum(1 for c in companies if c['status']=='ACTIVE')})")
    print(f"index レコード: {len(index_rows)}")
    print(f"要確認ページ: {len(candidates)}")
    print(f"前回 DeliveryLog: {len(previous_log)} 行継承")

    added_cids = sorted({r["company_code"] for r in index_rows})
    added_range = f"{added_cids[0]}-{added_cids[-1]}" if added_cids else ""
    new_log_row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        args.package_id, args.type, added_range,
        f"{len(index_rows)}docs", args.fde_name, args.note,
    ]

    wb = openpyxl.Workbook()
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    write_dashboard_sheet(ws_dash, companies, index_rows)
    write_documents_sheet(wb.create_sheet("Documents"), index_rows)
    write_candidates_sheet(wb.create_sheet("Candidates"), candidates)
    write_inactive_sheet(wb.create_sheet("Inactive"), companies)
    write_delivery_log_sheet(wb.create_sheet("DeliveryLog"), previous_log, new_log_row)
    write_config_sheet(wb.create_sheet("Config"), companies, index_rows)

    DASHBOARD_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DASHBOARD_XLSX)
    print(f"\n✓ {DASHBOARD_XLSX}")
    print(f"  package_id: {args.package_id}  type: {args.type}  note: {args.note}")


if __name__ == "__main__":
    main()
