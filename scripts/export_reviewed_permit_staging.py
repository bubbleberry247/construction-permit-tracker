#!/usr/bin/env python3
"""承認済みレビュー台帳からGAS取込用staging CSVを決定論的に出力する。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PERMIT_SHEET = "許可取込候補"
MONITORING_SHEET = "監視対象"
RECONCILIATION_SHEET = "許可照合"

PERMIT_HEADERS = [
    "migration_row_id",
    "source_sha256",
    "stable_key",
    "classification",
    "review_status",
    "reviewer",
    "reviewed_at",
    "review_note",
    "permit_id",
    "company_id",
    "company_name_raw",
    "permit_authority_name",
    "permit_authority_name_normalized",
    "permit_authority_type",
    "permit_category",
    "permit_year",
    "contractor_number",
    "permit_number_full",
    "trade_categories",
    "issue_date",
    "expiry_date",
    "renewal_deadline_date",
    "current_status",
    "mlit_confirmed_date",
    "mlit_confirm_result",
    "mlit_screenshot_url",
    "source_file",
    "source_file_hash",
    "parse_status",
    "note",
    "permit_data_version",
]

MONITORING_HEADERS = [
    "company_id",
    "vendor_no",
    "company_name_raw",
    "status",
    "has_permit",
    "has_mlit_observation",
    "proposed_monitoring_enabled",
    "proposed_action",
    "review_status",
    "reviewer",
    "reviewed_at",
    "review_note",
]

RECONCILIATION_HEADERS = [
    "reconciliation_id",
    "stable_key",
    "company_id",
    "company_name",
    "company_status",
    "authority",
    "contractor_number",
    "permit_category",
    "permit_count",
    "mlit_count",
    "permit_ids",
    "approved_expiry_date",
    "observed_expiry_date",
    "classification",
    "proposed_action",
    "risk_flags",
    "review_status",
    "reviewer",
    "reviewed_at",
    "review_note",
]


class ReviewValidationError(ValueError):
    """レビュー未完了または台帳形式不正。"""


def _normalized_header(value: Any) -> str:
    return str(value or "").lstrip("\ufeff").strip()


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_table(workbook: Any, sheet_name: str, expected_headers: list[str]) -> list[dict[str, str]]:
    if sheet_name not in workbook.sheetnames:
        raise ReviewValidationError(f"必要なsheetがありません: {sheet_name}")
    sheet = workbook[sheet_name]
    actual_headers = [
        _normalized_header(sheet.cell(1, column).value)
        for column in range(1, len(expected_headers) + 1)
    ]
    if actual_headers != expected_headers:
        raise ReviewValidationError(
            f"{sheet_name}のheaderが不正です: expected={expected_headers} actual={actual_headers}"
        )

    records: list[dict[str, str]] = []
    for row_number in range(2, sheet.max_row + 1):
        values = [
            _csv_value(sheet.cell(row_number, column).value)
            for column in range(1, len(expected_headers) + 1)
        ]
        if not any(values):
            continue
        if not values[0]:
            raise ReviewValidationError(f"{sheet_name}!A{row_number}の識別子が空です")
        record = dict(zip(expected_headers, values))
        record["_row"] = str(row_number)
        records.append(record)
    return records


def _require_manual_review_fields(record: dict[str, str], identifier: str) -> None:
    missing = [
        field
        for field in ("reviewer", "reviewed_at", "review_note")
        if not record.get(field, "").strip()
    ]
    if missing:
        raise ReviewValidationError(
            f"{identifier}は手動判断の証跡が不足しています: {','.join(missing)}"
        )


def _validate_reconciliation(records: list[dict[str, str]]) -> None:
    for record in records:
        classification = record["classification"].upper()
        if classification not in {"CONFLICT", "DUPLICATE"}:
            continue
        status = record["review_status"].upper()
        identifier = record["reconciliation_id"]
        if status not in {"APPROVED", "REJECTED"}:
            raise ReviewValidationError(
                f"{identifier}の{classification}が未解決です: {status or 'EMPTY'}"
            )
        _require_manual_review_fields(record, identifier)


def _validate_permits(records: list[dict[str, str]], expected_count: int) -> None:
    if len(records) != expected_count:
        raise ReviewValidationError(
            f"許可取込候補は{expected_count}件必要です: actual={len(records)}"
        )
    seen: set[str] = set()
    for record in records:
        identifier = record["migration_row_id"]
        if identifier in seen:
            raise ReviewValidationError(f"migration_row_idが重複しています: {identifier}")
        seen.add(identifier)
        status = record["review_status"].upper()
        if status not in {"APPROVED", "REJECTED"}:
            raise ReviewValidationError(
                f"{identifier}のreview_statusが未判断です: {status or 'EMPTY'}"
            )
        _require_manual_review_fields(record, identifier)
        if status == "APPROVED":
            missing = [
                field
                for field in (
                    "company_id",
                    "permit_authority_name",
                    "permit_category",
                    "contractor_number",
                    "expiry_date",
                )
                if not record.get(field, "").strip()
            ]
            if missing:
                raise ReviewValidationError(
                    f"{identifier}はAPPROVED必須項目が不足しています: {','.join(missing)}"
                )


def _validate_monitoring(records: list[dict[str, str]], expected_count: int) -> None:
    if len(records) != expected_count:
        raise ReviewValidationError(
            f"監視対象は{expected_count}社必要です: actual={len(records)}"
        )
    seen: set[str] = set()
    for record in records:
        identifier = record["company_id"]
        if identifier in seen:
            raise ReviewValidationError(f"company_idが重複しています: {identifier}")
        seen.add(identifier)
        status = record["review_status"].upper()
        action = record["proposed_action"].upper()
        if status not in {"APPROVED", "REJECTED", "AUTO_APPROVED"}:
            raise ReviewValidationError(
                f"{identifier}のreview_statusが未判断です: {status or 'EMPTY'}"
            )
        if action not in {"MONITOR", "DO_NOT_MONITOR"}:
            raise ReviewValidationError(f"{identifier}のproposed_actionが不正です: {action}")
        if status in {"APPROVED", "REJECTED"}:
            _require_manual_review_fields(record, identifier)
        elif not record["review_note"]:
            raise ReviewValidationError(f"{identifier}のAUTO_APPROVED理由がありません")
        if record["status"].upper() != "ACTIVE" and action == "MONITOR":
            raise ReviewValidationError(f"{identifier}はINACTIVEのためMONITORにできません")


def _write_csv(path: Path, headers: list[str], records: list[dict[str, str]]) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        for record in records:
            writer.writerow({header: record.get(header, "") for header in headers})
    text = path.read_text(encoding="utf-8-sig")
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def export_reviewed_staging(
    workbook_path: Path,
    output_dir: Path,
    *,
    expected_permit_count: int = 50,
    expected_monitoring_count: int = 131,
) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    reconciliation = _read_table(
        workbook, RECONCILIATION_SHEET, RECONCILIATION_HEADERS
    )
    permits = _read_table(workbook, PERMIT_SHEET, PERMIT_HEADERS)
    monitoring = _read_table(workbook, MONITORING_SHEET, MONITORING_HEADERS)

    _validate_reconciliation(reconciliation)
    _validate_permits(permits, expected_permit_count)
    _validate_monitoring(monitoring, expected_monitoring_count)

    permit_path = output_dir / "permit_import_staging.reviewed.csv"
    monitoring_path = output_dir / "monitoring_target_staging.reviewed.csv"
    permit_sha256 = _write_csv(permit_path, PERMIT_HEADERS, permits)
    monitoring_sha256 = _write_csv(monitoring_path, MONITORING_HEADERS, monitoring)
    workbook_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()

    manifest = {
        "workbook": str(workbook_path.resolve()),
        "workbook_sha256": workbook_sha256,
        "permit": {
            "path": str(permit_path.resolve()),
            "rows": len(permits),
            "sha256_normalized_text": permit_sha256,
            "load_confirmation": f"LOAD_PERMIT_STAGING_{permit_sha256[:12]}",
        },
        "monitoring": {
            "path": str(monitoring_path.resolve()),
            "rows": len(monitoring),
            "sha256_normalized_text": monitoring_sha256,
            "load_confirmation": f"LOAD_MONITORING_STAGING_{monitoring_sha256[:12]}",
        },
    }
    manifest_path = output_dir / "reviewed_staging_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    manifest["manifest_path"] = str(manifest_path.resolve())
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(
        description="承認済み許可レビュー台帳をGAS staging CSVへ変換します"
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    try:
        manifest = export_reviewed_staging(args.workbook, args.output_dir)
    except (OSError, ReviewValidationError) as error:
        parser.error(str(error))
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
