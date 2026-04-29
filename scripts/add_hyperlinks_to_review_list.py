"""
一覧表 v2 の各書類列「○」セルに、対応する PDF へのハイパーリンクと tooltip を付与する。

ZIP 解凍後の相対パス (10_originals/C0XXX_<会社名>/<カテゴリ>/<受領年>/<ファイル名>) で
リンクを設定。Excel で○をクリックするとローカル PDF ビューアで開く。
マウスオーバーすると相対パスが tooltip 表示される。

旧 DB の files.new_path には data/all_documents/ 配下のものが多く、
data/originals/ 構造に変換できないため、本スクリプトでは
**FDE_MANAGED/10_originals/ ディレクトリを直接スキャン** して
company_id + カテゴリフォルダから対応 PDF を見つける方針を採用。

Usage:
  python scripts/add_hyperlinks_to_review_list.py <xlsx_path>
  python scripts/add_hyperlinks_to_review_list.py output/FDE_MANAGED/00_一覧表_v2_20260427.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.hyperlink import Hyperlink

sys.stdout.reconfigure(encoding="utf-8")

PROJECT = Path(__file__).resolve().parent.parent

REQUIRED_DOCS = [
    "取引申請書", "建設業許可証", "決算書", "工事経歴書",
    "取引先一覧表", "労働安全衛生誓約書", "資格略字一覧", "労働者名簿",
]

CATEGORY_FOLDER = {
    "取引申請書":        "01_取引申請書",
    "建設業許可証":      "02_建設業許可証",
    "決算書":            "03_決算書",
    "工事経歴書":        "04_工事経歴書",
    "労働安全衛生誓約書": "05_労働安全衛生誓約書",
    "資格略字一覧":      "06_資格略字一覧",
    "取引先一覧表":      "07_取引先一覧表",
    "労働者名簿":        "08_労働者名簿",
}
BUNDLE_FOLDER = "99_受領バンドル"


def find_company_folder(originals_root: Path, company_id: str) -> Path | None:
    cands = [p for p in originals_root.glob(f"{company_id}_*") if p.is_dir()]
    if not cands:
        return None
    return cands[0]


def get_pdf_paths(originals_root: Path, company_id: str, doc_type: str) -> list[str]:
    """FDE_MANAGED/10_originals/C0XXX_*/<カテゴリ>/<年>/*.pdf + バンドル を相対パスで返す"""
    cdir = find_company_folder(originals_root, company_id)
    if cdir is None:
        return []
    paths: list[Path] = []
    cat_folder = CATEGORY_FOLDER.get(doc_type)
    if cat_folder:
        target = cdir / cat_folder
        if target.exists():
            paths.extend(sorted(p for p in target.rglob("*.pdf") if p.is_file()))
    # バンドルも候補 (1 PDF に複数書類が混在しているケース)
    bundle = cdir / BUNDLE_FOLDER
    if bundle.exists():
        for p in sorted(bundle.rglob("*.pdf")):
            if p.is_file() and p not in paths:
                paths.append(p)
    rels = []
    for p in paths:
        rel = p.relative_to(originals_root.parent).as_posix()  # 10_originals/C0XXX_*/...
        rels.append(rel)
    return rels


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx_path", type=Path)
    args = ap.parse_args()

    if not args.xlsx_path.exists():
        print(f"ERROR: {args.xlsx_path} not found", file=sys.stderr)
        sys.exit(1)

    # 同階層の 10_originals/ をスキャン (xlsx は FDE_MANAGED/ 直下にある前提)
    originals_root = args.xlsx_path.parent / "10_originals"
    if not originals_root.exists():
        print(f"ERROR: {originals_root} not found", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(args.xlsx_path)
    ws = wb["マスタ145社受領状況"]
    hdr = [c.value for c in ws[1]]
    cid_col = hdr.index("会社ID") + 1

    doc_cols = {}
    for d in REQUIRED_DOCS:
        if d in hdr:
            doc_cols[d] = hdr.index(d) + 1

    link_font = Font(color="0563C1", underline="single")
    n_links = 0
    n_skipped = 0
    skipped_samples = []

    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row, cid_col).value
        if not cid:
            continue
        for doc, col in doc_cols.items():
            cell = ws.cell(row, col)
            if cell.value != "○":
                continue
            paths = get_pdf_paths(originals_root, cid, doc)
            if not paths:
                n_skipped += 1
                if len(skipped_samples) < 10:
                    skipped_samples.append(f"{cid} / {doc}")
                continue
            primary = paths[0]
            if len(paths) == 1:
                tooltip = primary
            else:
                others = "\n  ".join(paths[1:])
                tooltip = f"{primary}\n他 {len(paths)-1} 件:\n  {others}"
            cell.hyperlink = Hyperlink(
                ref=cell.coordinate,
                target=primary,
                tooltip=tooltip,
                display="○",
            )
            cell.font = link_font
            n_links += 1

    wb.save(args.xlsx_path)

    print(f"=== ハイパーリンク追加完了 ===")
    print(f"  対象ファイル: {args.xlsx_path}")
    print(f"  リンク追加: {n_links} セル")
    print(f"  ファイル不在で skip: {n_skipped} セル")
    if skipped_samples:
        print(f"  skip サンプル (top 10):")
        for s in skipped_samples:
            print(f"    {s}")


if __name__ == "__main__":
    main()
