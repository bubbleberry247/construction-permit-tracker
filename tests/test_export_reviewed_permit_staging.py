from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from export_reviewed_permit_staging import (  # noqa: E402
    MONITORING_HEADERS,
    MONITORING_SHEET,
    PERMIT_HEADERS,
    PERMIT_SHEET,
    RECONCILIATION_HEADERS,
    RECONCILIATION_SHEET,
    ReviewValidationError,
    export_reviewed_staging,
)


def _append_dict(sheet, headers: list[str], values: dict[str, object]) -> None:
    sheet.append([values.get(header, "") for header in headers])


def _make_workbook(path: Path, *, conflict_status: str = "APPROVED", pending_permit: bool = False) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    reconciliation = workbook.create_sheet(RECONCILIATION_SHEET)
    reconciliation.append(RECONCILIATION_HEADERS)
    _append_dict(
        reconciliation,
        RECONCILIATION_HEADERS,
        {
            "reconciliation_id": "PR-0001",
            "classification": "CONFLICT",
            "review_status": conflict_status,
            "reviewer": "ops@example.com" if conflict_status != "PENDING" else "",
            "reviewed_at": "2026-07-30T15:00:00+09:00" if conflict_status != "PENDING" else "",
            "review_note": "証跡確認済み" if conflict_status != "PENDING" else "",
        },
    )

    permits = workbook.create_sheet(PERMIT_SHEET)
    permits.append(PERMIT_HEADERS)
    for index in range(2):
        _append_dict(
            permits,
            PERMIT_HEADERS,
            {
                "migration_row_id": f"PM-{index + 1:04d}",
                "source_sha256": "a" * 64,
                "stable_key": f"C{index + 1:04d}|愛知県知事|{index + 100}|一般",
                "classification": "MLIT_ONLY",
                "review_status": "PENDING" if pending_permit and index == 0 else "APPROVED",
                "reviewer": "ops@example.com",
                "reviewed_at": "2026-07-30T15:00:00+09:00",
                "review_note": "証跡確認済み",
                "company_id": f"C{index + 1:04d}",
                "permit_authority_name": "愛知県知事",
                "permit_category": "一般",
                "contractor_number": str(index + 100),
                "expiry_date": "2030-01-01",
                "permit_data_version": 1,
            },
        )

    monitoring = workbook.create_sheet(MONITORING_SHEET)
    monitoring.append(MONITORING_HEADERS)
    for index in range(3):
        _append_dict(
            monitoring,
            MONITORING_HEADERS,
            {
                "company_id": f"C{index + 1:04d}",
                "company_name_raw": f"会社{index + 1}",
                "status": "ACTIVE",
                "has_permit": index < 2,
                "has_mlit_observation": index < 2,
                "proposed_monitoring_enabled": index < 2,
                "proposed_action": "MONITOR" if index < 2 else "DO_NOT_MONITOR",
                "review_status": "APPROVED" if index < 2 else "AUTO_APPROVED",
                "reviewer": "ops@example.com" if index < 2 else "",
                "reviewed_at": "2026-07-30T15:00:00+09:00" if index < 2 else "",
                "review_note": "証跡確認済み" if index < 2 else "参照なしのため自動除外",
            },
        )
    workbook.save(path)


def test_reviewed_workbook_exports_deterministic_staging_csv(tmp_path: Path) -> None:
    workbook_path = tmp_path / "review.xlsx"
    _make_workbook(workbook_path)
    result = export_reviewed_staging(
        workbook_path,
        tmp_path / "out",
        expected_permit_count=2,
        expected_monitoring_count=3,
    )

    assert result["permit"]["rows"] == 2
    assert result["monitoring"]["rows"] == 3
    assert result["permit"]["load_confirmation"].startswith("LOAD_PERMIT_STAGING_")
    assert result["monitoring"]["load_confirmation"].startswith(
        "LOAD_MONITORING_STAGING_"
    )
    with Path(result["permit"]["path"]).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["review_status"] == "APPROVED"
    assert rows[0]["permit_data_version"] == "1"


def test_unresolved_conflict_is_rejected(tmp_path: Path) -> None:
    workbook_path = tmp_path / "review.xlsx"
    _make_workbook(workbook_path, conflict_status="PENDING")
    with pytest.raises(ReviewValidationError, match="CONFLICTが未解決"):
        export_reviewed_staging(
            workbook_path,
            tmp_path / "out",
            expected_permit_count=2,
            expected_monitoring_count=3,
        )


def test_pending_permit_is_rejected(tmp_path: Path) -> None:
    workbook_path = tmp_path / "review.xlsx"
    _make_workbook(workbook_path, pending_permit=True)
    with pytest.raises(ReviewValidationError, match="review_statusが未判断"):
        export_reviewed_staging(
            workbook_path,
            tmp_path / "out",
            expected_permit_count=2,
            expected_monitoring_count=3,
        )
