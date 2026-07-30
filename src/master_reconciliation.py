"""会社マスタの非破壊照合dry-run。

Excel原本と、Google Sheetsから読み取り専用で取得したJSON snapshotまたは
Excel exportを比較し、
MasterImportStaging投入前のCSVとsummary JSONを生成する。Google Sheetsは更新しない。

JSON snapshotはレコード配列、または ``{"records": [...]}`` を受け付ける。
Permits / MLITPermitsのレコードには company_id と会社名列のいずれかを含める。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import tempfile
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


STAGING_HEADERS = [
    "source_row",
    "vendor_no",
    "company_name_raw",
    "company_name_normalized",
    "matched_company_id",
    "classification",
    "match_method",
    "review_status",
    "reviewed_by",
    "reviewed_at",
    "notes",
    "source_sha256",
    "imported_at",
]


def normalize_company_name(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    text = text.replace("㈱", "株式会社").replace("㈲", "有限会社")
    text = re.sub(r"\s*[（(]株[）)]\s*", "株式会社", text)
    text = re.sub(r"\s*[（(]有[）)]\s*", "有限会社", text)
    return re.sub(r"\s+", "", text)


def normalize_vendor_no(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    if text in {"", "-", "－", "―"}:
        return ""
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_json_records(path: Path | None) -> list[dict[str, Any]]:
    if path is None:
        return []
    with path.open(encoding="utf-8-sig") as source:
        payload = json.load(source)
    if isinstance(payload, dict):
        payload = payload.get("records", payload.get("items", []))
    if not isinstance(payload, list):
        raise ValueError(f"JSON snapshot must be a list: {path}")
    return [dict(item) for item in payload if isinstance(item, dict)]


def values_to_records(values: list[list[Any]]) -> list[dict[str, Any]]:
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    records: list[dict[str, Any]] = []
    for row in values[1:]:
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        if not any(str(value or "").strip() for value in padded):
            continue
        records.append(
            {header: padded[index] for index, header in enumerate(headers) if header}
        )
    return records


def fetch_google_sheet_snapshots(
    service_account_file: Path,
    spreadsheet_id: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build

    credentials = Credentials.from_service_account_file(
        str(service_account_file),
        scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"],
    )
    service = build("sheets", "v4", credentials=credentials, cache_discovery=False)
    ranges = ["Companies!A1:Z1000", "Permits!A1:AG1000", "MLITPermits!A1:M200"]
    response = (
        service.spreadsheets()
        .values()
        .batchGet(
            spreadsheetId=spreadsheet_id,
            ranges=ranges,
            majorDimension="ROWS",
            valueRenderOption="UNFORMATTED_VALUE",
        )
        .execute()
    )
    value_ranges = response.get("valueRanges", [])
    if len(value_ranges) != 3:
        raise RuntimeError("Google Sheets snapshot response is incomplete")
    return tuple(
        values_to_records(value_range.get("values", []))
        for value_range in value_ranges
    )  # type: ignore[return-value]


def load_excel_rows(path: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency declared in requirements
        raise RuntimeError("openpyxl is required") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError("Excel source is empty")
    headers = [str(value or "").strip() for value in values[0]]

    def find_column(candidates: Iterable[str]) -> int:
        for candidate in candidates:
            for index, header in enumerate(headers):
                if header == candidate or candidate in header:
                    return index
        raise ValueError(f"required Excel column not found: {list(candidates)}")

    row_index = find_column(["№", "No", "NO"])
    vendor_index = find_column(["業者№", "業者No", "業者番号"])
    company_index = find_column(["会社名"])
    rows: list[dict[str, Any]] = []
    for excel_row, values_row in enumerate(values[1:], start=2):
        company_name = str(values_row[company_index] or "").strip()
        if not company_name:
            continue
        rows.append(
            {
                "source_row": excel_row,
                "source_no": values_row[row_index],
                "vendor_no": normalize_vendor_no(values_row[vendor_index]),
                "company_name_raw": company_name,
                "company_name_normalized": normalize_company_name(company_name),
            }
        )
    return rows


def load_live_workbook_snapshots(
    path: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Google Sheetsから読み取り専用でexportしたxlsxをレコードへ変換する。"""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency declared in requirements
        raise RuntimeError("openpyxl is required") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def read_sheet(name: str) -> list[dict[str, Any]]:
        if name not in workbook.sheetnames:
            raise ValueError(f"required live workbook sheet not found: {name}")
        values = [list(row) for row in workbook[name].iter_rows(values_only=True)]
        return values_to_records(values)

    return read_sheet("Companies"), read_sheet("Permits"), read_sheet("MLITPermits")


def first_nonempty(record: dict[str, Any], fields: Iterable[str]) -> str:
    for field in fields:
        value = str(record.get(field) or "").strip()
        if value:
            return value
    return ""


@dataclass
class SystemCompany:
    company_id: str
    names: set[str]
    raw_names: set[str]
    sources: set[str]


def build_system_companies(
    companies: list[dict[str, Any]],
    permits: list[dict[str, Any]],
    mlit_permits: list[dict[str, Any]],
) -> dict[str, SystemCompany]:
    combined: dict[str, SystemCompany] = {}

    def add(record: dict[str, Any], source: str, name_fields: list[str]) -> None:
        company_id = str(record.get("company_id") or "").strip()
        if not company_id:
            return
        raw_name = first_nonempty(record, name_fields)
        company = combined.setdefault(
            company_id,
            SystemCompany(company_id=company_id, names=set(), raw_names=set(), sources=set()),
        )
        company.sources.add(source)
        if raw_name:
            company.raw_names.add(raw_name)
            company.names.add(normalize_company_name(raw_name))

    for row in companies:
        add(row, "Companies", ["company_name_raw", "company_name_normalized", "company_name"])
    for row in permits:
        add(row, "Permits", ["company_name_raw", "company_name", "company_name_normalized"])
    for row in mlit_permits:
        add(row, "MLITPermits", ["company_name", "company_name_raw", "company_name_normalized"])
    return combined


def reconcile(
    excel_rows: list[dict[str, Any]],
    system_companies: dict[str, SystemCompany],
    source_hash: str,
    imported_at: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    raw_index: dict[str, set[str]] = defaultdict(set)
    normalized_index: dict[str, set[str]] = defaultdict(set)
    for company_id, company in system_companies.items():
        for raw_name in company.raw_names:
            raw_index[raw_name].add(company_id)
        for normalized in company.names:
            if normalized:
                normalized_index[normalized].add(company_id)

    excel_raw_counts = Counter(row["company_name_raw"] for row in excel_rows)
    vendor_counts = Counter(row["vendor_no"] for row in excel_rows if row["vendor_no"])
    staging: list[dict[str, Any]] = []
    matched_ids: set[str] = set()

    for row in excel_rows:
        raw_name = row["company_name_raw"]
        normalized = row["company_name_normalized"]
        vendor_no = row["vendor_no"]
        exact_ids = raw_index.get(raw_name, set())
        normalized_ids = normalized_index.get(normalized, set())
        classification = "NEW_FROM_EXCEL"
        method = "NONE"
        matched_id = ""
        notes: list[str] = []

        if excel_raw_counts[raw_name] > 1 or (vendor_no and vendor_counts[vendor_no] > 1):
            classification = "DUPLICATE"
            method = "SOURCE_DUPLICATE"
            notes.append("Excel内で会社名または業者番号が重複")
        elif len(exact_ids) == 1:
            classification = "MATCHED"
            method = "EXACT_RAW_NAME"
            matched_id = next(iter(exact_ids))
        elif len(exact_ids) > 1:
            classification = "CONFLICT"
            method = "EXACT_RAW_NAME_MULTIPLE"
            notes.append("同一表記に複数company_id")
        elif len(normalized_ids) == 1:
            classification = "REVIEW_REQUIRED"
            method = "NORMALIZED_NAME_CANDIDATE"
            matched_id = next(iter(normalized_ids))
            notes.append("正規化名のみ一致。原表記を人手確認")
        elif len(normalized_ids) > 1:
            classification = "CONFLICT"
            method = "NORMALIZED_NAME_MULTIPLE"
            notes.append("正規化名に複数company_id")

        if matched_id:
            matched_ids.add(matched_id)
        staging.append(
            {
                "source_row": row["source_row"],
                "vendor_no": vendor_no,
                "company_name_raw": raw_name,
                "company_name_normalized": normalized,
                "matched_company_id": matched_id,
                "classification": classification,
                "match_method": method,
                "review_status": "PENDING",
                "reviewed_by": "",
                "reviewed_at": "",
                "notes": "; ".join(notes),
                "source_sha256": source_hash,
                "imported_at": imported_at,
            }
        )

    system_only: list[dict[str, Any]] = []
    for company_id, company in sorted(system_companies.items()):
        if company_id in matched_ids:
            continue
        system_only.append(
            {
                "company_id": company_id,
                "company_names": " / ".join(sorted(company.raw_names)),
                "sources": ",".join(sorted(company.sources)),
                "classification": "SYSTEM_ONLY",
                "review_status": "PENDING",
            }
        )

    classification_counts = Counter(row["classification"] for row in staging)
    orphan_ids = sorted(
        company_id
        for company_id, company in system_companies.items()
        if "Permits" in company.sources or "MLITPermits" in company.sources
        if not company.raw_names
    )
    summary = {
        "source_sha256": source_hash,
        "imported_at": imported_at,
        "excel_company_count": len(excel_rows),
        "excel_vendor_no_count": sum(1 for row in excel_rows if row["vendor_no"]),
        "excel_vendor_no_blank_count": sum(1 for row in excel_rows if not row["vendor_no"]),
        "system_company_id_count": len(system_companies),
        "classification_counts": dict(sorted(classification_counts.items())),
        "system_only_count": len(system_only),
        "permit_reference_without_name_count": len(orphan_ids),
        "permit_reference_without_name_ids": orphan_ids,
        "approval_ready": (
            len(excel_rows) == 127
            and classification_counts.get("DUPLICATE", 0) == 0
            and classification_counts.get("CONFLICT", 0) == 0
            and classification_counts.get("REVIEW_REQUIRED", 0) == 0
            and len(system_only) == 0
            and len(orphan_ids) == 0
        ),
    }
    return staging, system_only, summary


def atomic_write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=path.stem + "_", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as target:
            json.dump(
                payload,
                target,
                ensure_ascii=False,
                indent=2,
                default=json_default,
            )
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def json_default(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, set):
        return sorted(value)
    raise TypeError(f"Object of type {type(value).__name__} is not JSON serializable")


def atomic_write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=path.stem + "_", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8-sig", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--companies-json", type=Path)
    parser.add_argument("--permits-json", type=Path)
    parser.add_argument("--mlit-json", type=Path)
    parser.add_argument(
        "--live-xlsx",
        type=Path,
        help="Google Sheetsの読み取り専用Excel export",
    )
    parser.add_argument("--spreadsheet-id")
    parser.add_argument("--service-account-file", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--expected-count", type=int, default=127)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source_hash = sha256_file(args.xlsx)
    excel_rows = load_excel_rows(args.xlsx)
    if len(excel_rows) != args.expected_count:
        raise ValueError(
            f"Excel company count mismatch: expected={args.expected_count} actual={len(excel_rows)}"
        )
    source_modes = sum(
        [
            bool(args.live_xlsx),
            bool(args.spreadsheet_id or args.service_account_file),
            bool(args.companies_json or args.permits_json or args.mlit_json),
        ]
    )
    if source_modes > 1:
        raise ValueError(
            "--live-xlsx, Google Sheets API, JSON snapshotsは同時指定できません"
        )

    snapshot_source = ""
    if args.live_xlsx:
        companies, permits, mlit_permits = load_live_workbook_snapshots(args.live_xlsx)
        snapshot_source = str(args.live_xlsx)
    elif args.spreadsheet_id or args.service_account_file:
        if not args.spreadsheet_id or not args.service_account_file:
            raise ValueError(
                "--spreadsheet-id and --service-account-file must be specified together"
            )
        companies, permits, mlit_permits = fetch_google_sheet_snapshots(
            args.service_account_file, args.spreadsheet_id
        )
        snapshot_source = args.spreadsheet_id
    else:
        companies = load_json_records(args.companies_json)
        permits = load_json_records(args.permits_json)
        mlit_permits = load_json_records(args.mlit_json)
        snapshot_source = "JSON"
    system_companies = build_system_companies(companies, permits, mlit_permits)
    imported_at = datetime.now().astimezone().isoformat(timespec="seconds")
    staging, system_only, summary = reconcile(
        excel_rows, system_companies, source_hash, imported_at
    )
    summary["live_snapshot_source"] = snapshot_source
    summary["live_snapshot_sha256"] = (
        sha256_file(args.live_xlsx) if args.live_xlsx else ""
    )
    atomic_write_csv(args.output_dir / "master_import_staging.csv", STAGING_HEADERS, staging)
    atomic_write_csv(
        args.output_dir / "system_only.csv",
        ["company_id", "company_names", "sources", "classification", "review_status"],
        system_only,
    )
    atomic_write_json(args.output_dir / "summary.json", summary)
    atomic_write_json(
        args.output_dir / "snapshot_companies.json",
        {"snapshot_source": snapshot_source, "records": companies},
    )
    atomic_write_json(
        args.output_dir / "snapshot_permits.json",
        {"snapshot_source": snapshot_source, "records": permits},
    )
    atomic_write_json(
        args.output_dir / "snapshot_mlit_permits.json",
        {"snapshot_source": snapshot_source, "records": mlit_permits},
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
